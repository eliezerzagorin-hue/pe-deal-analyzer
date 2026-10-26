"""Fictional template tests; no third-party deal information is stored."""
from io import BytesIO
import pytest
from openpyxl import Workbook, load_workbook
from deallens.template_mapper import PROFILE, inspect_template, populate_template

def fake_template():
    wb = Workbook()
    wb.active.title = "ASSUMPTIONS"
    for name in ["REV", "EXP", "SUMMARY", "DEBT", "IS", "R"]:
        wb.create_sheet(name)
    s = wb["ASSUMPTIONS"]
    for key, _name, sheet, address, marker, _lo, _hi in PROFILE:
        row = s[address].row
        s[f"B{row}"] = marker
        s[address] = 2 if key == "case" else 1.0
    s["F17"] = "=IS!H31"
    s["F19"] = "=F18*F17"
    s["F27"] = "=F25+F26"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def test_populate_does_not_overwrite_formulas():
    data = fake_template()
    preview = inspect_template(data)
    assert len(preview["inputs"]) == len(PROFILE)
    output, addresses = populate_template(data, {"entry_multiple": 7.5, "exit_multiple": 7.0})
    result = load_workbook(BytesIO(output), data_only=False)
    assert addresses == ["ASSUMPTIONS!F18", "ASSUMPTIONS!F20"]
    assert result["ASSUMPTIONS"]["F18"].value == 7.5
    assert result["ASSUMPTIONS"]["F20"].value == 7.0
    assert result["ASSUMPTIONS"]["F17"].value == "=IS!H31"
    assert result["ASSUMPTIONS"]["F19"].value == "=F18*F17"

def test_reject_formula_target():
    wb = load_workbook(BytesIO(fake_template()))
    wb["ASSUMPTIONS"]["F18"] = "=2+3"
    b = BytesIO()
    wb.save(b)
    with pytest.raises(ValueError, match="formula"):
        populate_template(b.getvalue(), {"entry_multiple": 8})

def test_reject_unknown_and_out_of_range():
    data = fake_template()
    with pytest.raises(ValueError, match="Unknown"):
        populate_template(data, {"fake_input": 3})
    with pytest.raises(ValueError, match="Out-of-range"):
        populate_template(data, {"exit_multiple": 900})
    with pytest.raises(ValueError, match="integer"):
        populate_template(data, {"case": 1.5})

def test_reject_wrong_layout():
    wb=Workbook()
    b=BytesIO()
    wb.save(b)
    with pytest.raises(ValueError, match="profile"):
        inspect_template(b.getvalue())
