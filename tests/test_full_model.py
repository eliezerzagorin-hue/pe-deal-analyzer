"""Fictional data only: test all 14 sheets, source comments, and recommendation gates."""
from io import BytesIO

import pytest
from openpyxl import Workbook,load_workbook
from deallens.full_model import REQUIRED,make_integrated_model,recommendation,read_recalculated_returns
from deallens.ic_brief import export_two_page_ic
from deallens.demo import demo_memo

def _template():
    wb=Workbook()
    wb.active.title=REQUIRED[0]
    for title in REQUIRED[1:]:wb.create_sheet(title)
    wb["COVER"]["D14"]="OLD COMPANY"
    wb["ASSUMPTIONS"]["F18"]=7.5
    wb["REV"]["E9"]=99999
    wb["EXP"]["E10"]=50000
    wb["COMPS"]["B10"]="Old comp"
    wb["COMPS"]["F10"]=100000
    wb["R"]["E21"]="=1+1"
    b=BytesIO();wb.save(b);return b.getvalue()

def _annual():
    periods=[]
    for year in range(2022,2030):
        periods.append(dict(year=year,kind="actual" if year<=2025 else "projection",
             revenue=10+(year-2022),materials=2,labor=1,other_cogs=1,
             total_cogs=4,executive=.5,administration=.3,rent=.2,sales=.1,
             maintenance=.1,production=.1,other_opex=.4,total_opex=1.7,
             depreciation=.3,capex=.4,net_income=.8,interest=.1,income_tax=.2,
             adjusted_ebitda=2.6,source_page=4))
    return {"periods":periods,"warnings":[]}

def test_populate_all_14_tabs_with_cim_sources():
    data,issues=make_integrated_model(_template(),"Fictional Manufacturing",_annual(),base_year=2025)
    wb=load_workbook(BytesIO(data),data_only=False)
    assert wb.sheetnames==REQUIRED
    assert wb["COVER"]["D14"].value=="Fictional Manufacturing"
    assert wb["REV"]["E9"].value==10000
    assert wb["REV"]["I14"].value==14000
    assert "CIM PDF page 4" in wb["REV"]["I14"].comment.text
    assert wb["REV"]["M14"].data_type=="f"
    assert wb["COMPS"]["B10"].value is None
    assert wb["COMPS"]["F10"].value is None
    assert wb["R"]["E21"].data_type=="f"
    assert wb["ASSUMPTIONS"]["F18"].value==7.0
    assert wb["SUMMARY"]["B101"].value.startswith("SOURCE AUDIT")
    assert issues

def test_strict_return_gate():
    assert recommendation(3.0,.25,True,[])[0]=="CONDITIONAL PROCEED"
    assert recommendation(2.99,.25,True,[])[0]=="DO NOT PROCEED"
    assert recommendation(3.5,.249,True,[])[0]=="DO NOT PROCEED"
    assert recommendation(3.5,.30,False,[])[0]=="HOLD / FURTHER DILIGENCE"
    assert recommendation(3.5,.30,True,["Unreconciled cash flow"])[0]=="HOLD / FURTHER DILIGENCE"

def test_word_memo_has_all_sections_and_explicit_page_break():
    data=export_two_page_ic("Demo Company",demo_memo(),moic=2.9,irr=.24,verified=True,issues=[])
    doc=__import__("docx").Document(BytesIO(data))
    all_text=" ".join(p.text for p in doc.paragraphs)
    assert "COMPANY / BUSINESS OVERVIEW" in all_text
    assert "RECOMMENDATION" in all_text
    assert "DO NOT PROCEED" in all_text
    import zipfile
    with zipfile.ZipFile(BytesIO(data)) as z:
        xml=z.read("word/document.xml")
    assert b'w:type="page"' in xml
