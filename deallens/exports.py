"""Memo Word export and live-formula Excel LBO output."""
from io import BytesIO
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from .memo import FIELD_GROUPS

def export_memo(company, sections):
    doc=Document()
    doc.add_heading("INVESTMENT COMMITTEE MEMORANDUM",0)
    doc.add_paragraph(company)
    doc.add_paragraph("DRAFT. Requires source verification and investment-team review.")
    for title,fields in FIELD_GROUPS:
        doc.add_heading(title,1)
        for key,label in fields:
            if len(fields)>1:doc.add_heading(label,2)
            for text in sections.get(key,"Further diligence required.").splitlines():
                if text.strip():doc.add_paragraph(text.strip())
    doc.add_paragraph("AI citations and all numerical claims must be independently verified.")
    b=BytesIO();doc.save(b);return b.getvalue()

def export_lbo(a,results):
    wb=Workbook();inp=wb.active;inp.title="Inputs"
    vals=[("Entry revenue",a.entry_revenue),("Entry EBITDA",a.entry_ebitda),
      ("Entry multiple",a.entry_multiple),("Debt multiple",a.debt_multiple),
      ("Fees / EV",a.deal_fees_pct),("Revenue growth",a.revenue_growth),
      ("EBITDA margin",a.entry_margin),("Annual margin change",a.annual_margin_change),
      ("Interest rate",a.interest_rate),("Tax rate",a.tax_rate),
      ("D&A / sales",a.da_pct_revenue),("Capex / sales",a.capex_pct_revenue),
      ("NWC / incremental sales",a.nwc_pct_revenue),("Exit multiple",a.exit_multiple)]
    inp.append(["INPUTS | $M","Value"])
    for name,value in vals:inp.append([name,value])
    inp.column_dimensions["A"].width=32;inp.column_dimensions["B"].width=18
    for row in range(6,15):inp.cell(row,2).number_format="0.0%"
    ws=wb.create_sheet("LBO")
    rows=["Year","Revenue","EBITDA margin","EBITDA","D&A","EBIT","Opening debt","Interest expense",
          "Cash taxes","Capex","NWC investment","Free cash flow","Debt paydown",
          "Sponsor injection","Ending debt","Exit EV","Exit equity","Sponsor cash flow"]
    for i,v in enumerate(rows,1):ws.cell(i,1,v)
    ws.column_dimensions["A"].width=29
    for year in range(a.hold_years+1):
        col=year+2;c=get_column_letter(col);p=get_column_letter(col-1)
        ws.cell(1,col,year);ws.column_dimensions[c].width=18
        if year==0:
            forms={2:"=Inputs!B2",3:"=Inputs!B8",4:"=Inputs!B3",15:"=Inputs!B3*Inputs!B5",
              18:"=-(Inputs!B3*Inputs!B4*(1+Inputs!B6)-Inputs!B3*Inputs!B5)"}
        else:
            forms={2:f"={p}2*(1+Inputs!$B$7)",3:f"=Inputs!$B$8+{c}1*Inputs!$B$9",
              4:f"={c}2*{c}3",5:f"={c}2*Inputs!$B$12",6:f"={c}4-{c}5",
              7:f"={p}15",8:f"={c}7*Inputs!$B$10",
              9:f"=MAX(0,{c}6-{c}8)*Inputs!$B$11",
              10:f"={c}2*Inputs!$B$13",11:f"=({c}2-{p}2)*Inputs!$B$14",
              12:f"={c}4-{c}8-{c}9-{c}10-{c}11",
              13:f"=MIN({c}7,MAX(0,{c}12))",14:f"=MAX(0,-{c}12)",
              15:f"={c}7-{c}13",
              16:f"={c}4*Inputs!$B$15" if year==a.hold_years else "=0",
              17:f"=MAX(0,{c}16-{c}15)" if year==a.hold_years else "=0",
              18:f"={c}17-{c}14"}
        for row,formula in forms.items():ws.cell(row,col,formula)
        ws.cell(3,col).number_format="0.0%"
    last=get_column_letter(a.hold_years+2)
    out=wb.create_sheet("Returns")
    for row in [["Return summary","Formula"],["Entry EV","=Inputs!B3*Inputs!B4"],
       ["Initial sponsor equity","=-LBO!B18"],["Exit EV",f"=LBO!{last}16"],
       ["Exit debt",f"=LBO!{last}15"],["Exit equity",f"=LBO!{last}17"],
       ["Total capital contributed",f"=B3+SUM(LBO!C14:{last}14)"],
       ["MOIC","=IFERROR(B6/B7,0)"],["IRR",f"=IFERROR(IRR(LBO!B18:{last}18),0)"]]:out.append(row)
    out["B9"].number_format="0.0%";out.column_dimensions["A"].width=29;out.column_dimensions["B"].width=22
    for sheet in wb:
        sheet.freeze_panes="B2"
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="163858")
    ws["A20"]="Single-tranche model; opening-debt interest; annual sweep; no revolver."
    wb.calculation.fullCalcOnLoad=True
    b=BytesIO();wb.save(b);return b.getvalue()
