"""Compact decision-first IC memo: two deliberately composed Word pages."""
from io import BytesIO
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY=RGBColor(23,47,75)
TEAL=RGBColor(0,114,133)
GRAY=RGBColor(95,105,117)

def _tidy(value,limit=500):
    text=re.sub(r"\\s+"," ",str(value or "Further diligence required.")).strip()
    if len(text)<=limit: return text
    cut=text[:limit].rfind(". ")
    if cut<limit//2: cut=limit
    return text[:cut].rstrip(" ,;") + ("." if not text[:cut].endswith(".") else "")

def _title(doc,text):
    p=doc.add_paragraph()
    p.space_before=Pt(6);p.space_after=Pt(2)
    r=p.add_run(text.upper());r.bold=True;r.font.size=Pt(9);r.font.color.rgb=TEAL

def _body(doc,text,limit=470):
    p=doc.add_paragraph(_tidy(text,limit))
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.line_spacing=1.02

def _row(doc,left,right):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(2)
    r=p.add_run(left+": ");r.bold=True;r.font.size=Pt(8.3);r.font.color.rgb=NAVY
    r=p.add_run(_tidy(right,190));r.font.size=Pt(8.3)

def _cell(cell,text,bold=False):
    cell.text=text
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.space_before=Pt(0)
        for r in p.runs:
            r.font.size=Pt(8);r.bold=bold

def _header_footer(doc):
    sec=doc.sections[0]
    sec.top_margin=Inches(.54);sec.bottom_margin=Inches(.53)
    sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
    sec.header_distance=Inches(.2);sec.footer_distance=Inches(.24)
    h=sec.header.paragraphs[0]
    h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=h.add_run("DEALLENS  /  INTERNAL SCREENING DRAFT")
    r.font.size=Pt(7);r.bold=True;r.font.color.rgb=GRAY
    f=sec.footer.paragraphs[0]
    f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run("Preliminary. Verify financials and CIM citations.  •  Page ")
    r.font.size=Pt(7);r.font.color.rgb=GRAY
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
    f._p.append(field)

def export_two_page_ic(company,memo,moic=None,irr=None,verified=False,issues=None):
    """Precisely two planned pages. Text is capped to protect page count."""
    from .full_model import recommendation
    status,rationale=recommendation(moic,irr,verified,issues)
    doc=Document()
    style=doc.styles["Normal"];style.font.name="Aptos";style.font.size=Pt(8.5)
    style.font.color.rgb=NAVY
    style.paragraph_format.space_after=Pt(2)
    _header_footer(doc)
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(0)
    r=p.add_run("INVESTMENT COMMITTEE  /  INITIAL SCREEN")
    r.font.size=Pt(9);r.bold=True;r.font.color.rgb=TEAL
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(1)
    r=p.add_run(_tidy(company,110));r.bold=True;r.font.size=Pt(16);r.font.color.rgb=NAVY
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(8)
    r=p.add_run("Recommendation: "+status)
    r.bold=True;r.font.size=Pt(10)
    r.font.color.rgb=TEAL if status=="CONDITIONAL PROCEED" else RGBColor(165,64,43)
    t=doc.add_table(rows=2,cols=3)
    t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,head in enumerate(["Sponsor MOIC","Sponsor IRR","Minimum hurdles"]):_cell(t.cell(0,i),head,True)
    for i,value in enumerate([f"{moic:.2f}x" if moic is not None else "Pending model validation",
                              f"{irr:.1%}" if irr is not None else "Pending model validation",
                              "5-YR: 25% IRR / ~3.05x"]):_cell(t.cell(1,i),value)
    _title(doc,"1. Company / Business Overview")
    _row(doc,"Management & Employees",memo.get("management_employees",""))
    _row(doc,"Products & Services",memo.get("products_services",""))
    _row(doc,"General Operations",memo.get("general_operations",""))
    _row(doc,"Customer Base",memo.get("customer_base",""))
    _title(doc,"2. Historical Financial Performance")
    _body(doc,memo.get("historical_financial_performance",""),340)
    _title(doc,"3. Market / Industry Dynamics / Competition")
    _body(doc,memo.get("market_overview",""),370)
    _title(doc,"4. Growth Opportunities")
    _body(doc,memo.get("growth_opportunities",""),400)
    doc.add_page_break()
    p=doc.add_paragraph()
    r=p.add_run(_tidy(company,100)+"  |  IC Screening (continued)")
    r.bold=True;r.font.color.rgb=NAVY;r.font.size=Pt(11)
    _title(doc,"5. Investment Merits, Risks & Mitigants")
    _body(doc,memo.get("investment_merits_risks_mitigants",""),650)
    _title(doc,"6. Investment Thesis")
    _body(doc,memo.get("investment_thesis",""),510)
    _title(doc,"7. Potential Exit Strategies / Buyer Universe")
    _body(doc,memo.get("potential_exit_strategies_buyer_universe",""),510)
    _title(doc,"8. Recommendation")
    _body(doc,status+". "+rationale,350)
    _body(doc,"Memo narrative is not an authorization to transact. Only the independently verified sponsor-return gate above controls the recommendation.",200)
    _title(doc,"Open Diligence & Return Verification")
    if issues:
        _body(doc,"; ".join(issues[:4]),390)
    else:
        _body(doc,"Source claims, projections, adjusted EBITDA bridge, customer concentration, capital expenditures, financing terms and exit multiple still require independent diligence.",330)
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10)
    r=p.add_run("Decision policy: Five-year hold. Proceed only with verified >=25% sponsor IRR (equivalent to >=3.05x MOIC). Exactly 3.00x over five years is 24.6% IRR.")
    r.font.size=Pt(8);r.bold=True;r.font.color.rgb=TEAL
    b=BytesIO();doc.save(b);return b.getvalue()
