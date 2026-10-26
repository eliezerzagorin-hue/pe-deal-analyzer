"""Private, reusable LBO template loader and readable template diagnostics.

The confidential XLSX must never be committed to the public repository.
Configured Streamlit Secrets provide the persistent copy; session uploads
are temporary and isolated to the visitor's Streamlit session.
"""
import base64
import gzip
import io
import os
import zipfile

MAX_COMPRESSED_BYTES = 8_000_000
MAX_WORKBOOK_BYTES = 15_000_000
KEY = "LBO_TEMPLATE_GZIP_BASE64"
REQUIRED_SHEETS = (
    "COVER", "ASSUMPTIONS", "SUMMARY", "REV", "EXP", "D&A", "NWC",
    "DEBT", "IS", "BS", "CF", "R", "COMPS", "CAP",
)


def validate_integrated_template(data):
    """Check the user's 14-sheet workbook *before* CIM modeling starts.

    Return its sheet count and canonical names or raise a readable error.
    Keeping this check separate from CIM validation makes the UI clear about
    whether the file itself is recognized.
    """
    from openpyxl import load_workbook

    if not isinstance(data, bytes) or not data:
        raise ValueError("No Excel template has been supplied.")
    if len(data) > MAX_WORKBOOK_BYTES:
        raise ValueError("The Excel template is larger than 15 MB.")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False,
                                 keep_links=False)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise ValueError("The template could not be read as a valid .xlsx workbook.") from exc
    try:
        actual = tuple(workbook.sheetnames)
        if actual != REQUIRED_SHEETS:
            missing = [x for x in REQUIRED_SHEETS if x not in actual]
            extra = [x for x in actual if x not in REQUIRED_SHEETS]
            if not missing and not extra:
                raise ValueError(
                    "All 14 expected tabs exist, but their order differs from the "
                    "original model. The expected order is: " + ", ".join(REQUIRED_SHEETS)
                )
            raise ValueError(
                "This Excel file is not the original 14-tab model. "
                + (("Missing sheets: " + ", ".join(missing) + ". ") if missing else "")
                + (("Unexpected sheets: " + ", ".join(extra) + ". ") if extra else "")
                + "Required: " + ", ".join(REQUIRED_SHEETS)
            )
        if workbook["ASSUMPTIONS"]["F18"].data_type == "f":
            raise ValueError("ASSUMPTIONS!F18 is a formula, not the expected entry multiple input.")
        if workbook["REV"]["E9"].data_type == "f":
            raise ValueError("REV!E9 is a formula, not the expected historical revenue input.")
        return {"sheet_count": len(actual), "sheet_names": list(actual)}
    finally:
        workbook.close()


def decode_template(encoded):
    """Decode saved private template. Gzip/base64 is NOT encryption."""
    if not isinstance(encoded, str) or not encoded.strip():
        raise ValueError("The saved LBO template setting is empty.")
    encoded = "".join(encoded.split())
    if len(encoded) > MAX_COMPRESSED_BYTES * 2:
        raise ValueError("The saved LBO template setting is too large.")
    try:
        compressed = base64.b64decode(encoded, validate=True)
        if len(compressed) > MAX_COMPRESSED_BYTES:
            raise ValueError("Compressed LBO template exceeds the size limit.")
        with gzip.GzipFile(fileobj=io.BytesIO(compressed), mode="rb") as stream:
            content = stream.read(MAX_WORKBOOK_BYTES + 1)
    except (OSError, ValueError) as exc:
        raise ValueError(
            "The saved template was not decoded. Its gzip/base64 data may be truncated. "
            "Copy the entire .toml configuration into the Streamlit app's Secrets settings."
        ) from exc
    validate_integrated_template(content)
    return content


def get_saved_template(secrets=None):
    """Return an XLSX only when persistent private configuration is complete.

    Prefer TOML section keys, since large root-level values may exceed
    operating-system environment-variable limits.
    """
    data = None
    if secrets is not None:
        try:
            section = secrets.get("lbo_template", {})
            data = section.get("gzip_base64") if section else None
            if not data:
                data = secrets.get(KEY)
        except (KeyError, FileNotFoundError, RuntimeError):
            data = None
    if not data:
        data = os.environ.get(KEY)
    if not data:
        return None
    return decode_template(data)
