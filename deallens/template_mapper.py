"""Workbook-aware editable input mapping for a 14-sheet PE model.

This module contains only structural cell locations, never source deal data.
It deliberately refuses to overwrite formulas or export without explicit approval.
"""
from io import BytesIO
import math

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties

MAX_TEMPLATE_BYTES = 12 * 1024 * 1024

# (id, label, worksheet, cell, required existing label, min, max).
# Source units follow the uploaded workbook. Monetary fields are in USD THOUSANDS.
PROFILE = [
    ("case", "Operating case (1 management, 2 base, 3 downside)", "ASSUMPTIONS", "F9", "Active Case Selector", 1, 3),
    ("entry_multiple", "Entry EV / EBITDA (x)", "ASSUMPTIONS", "F18", "Entry EV / EBITDA Multiple", 0.01, 40),
    ("exit_multiple", "Exit EV / EBITDA (x)", "ASSUMPTIONS", "F20", "Exit EV / EBITDA Multiple", 0.01, 40),
    ("leverage", "Closing debt / EBITDA (x)", "ASSUMPTIONS", "F23", "Total Leverage at Close", 0, 20),
    ("sofr", "SOFR rate", "ASSUMPTIONS", "F25", "SOFR", 0, 1),
    ("term_spread", "Term loan credit spread", "ASSUMPTIONS", "F26", "Term Loan Credit Spread", 0, 1),
    ("revolver_cap", "Revolver capacity ($000s)", "ASSUMPTIONS", "F28", "Revolver Capacity", 0, 1e9),
    ("revolver_spread", "Revolver credit spread", "ASSUMPTIONS", "F29", "Revolver Credit Spread", 0, 1),
    ("amortization", "Mandatory annual amortization", "ASSUMPTIONS", "F31", "Mandatory Amortization", 0, 1),
    ("cash_sweep", "Cash-flow sweep", "ASSUMPTIONS", "F32", "Cash Flow Sweep", 0, 1),
    ("minimum_cash", "Minimum operating cash ($000s)", "ASSUMPTIONS", "F33", "Minimum Operating Cash", 0, 1e9),
    ("owner_rollover", "Owner rollover ($000s)", "ASSUMPTIONS", "F37", "Rollover Equity - Rod Fraser", 0, 1e9),
    ("management_roll", "Management rollover (% of old stake)", "ASSUMPTIONS", "F38", "Rollover Equity - Management", 0, 1),
    ("financing_fee", "Financing fees (% of debt)", "ASSUMPTIONS", "F41", "Financing Fees", 0, 1),
    ("tax_rate", "Tax rate", "ASSUMPTIONS", "F45", "Pro Forma Tax Rate", 0, 1),
    ("annual_capex", "Annual capex ($000s)", "ASSUMPTIONS", "F46", "Annual Capex", 0, 1e9),
    ("tax_stepup", "338(h)(10) step-up (0=no, 1=yes)", "ASSUMPTIONS", "F47", "338(h)(10)", 0, 1),
    ("tax_amort_years", "Tax amortization years", "ASSUMPTIONS", "F48", "Tax Amortization Period", 1, 50),
]

def _load(data):
    if not data or len(data) > MAX_TEMPLATE_BYTES:
        raise ValueError("Upload a nonempty .xlsx workbook under 12 MB.")
    try:
        return load_workbook(BytesIO(data), data_only=False, keep_links=True)
    except Exception as exc:
        raise ValueError("Could not open this Excel template.") from exc

def _validate_layout(wb):
    required = {"ASSUMPTIONS", "REV", "EXP", "SUMMARY", "DEBT", "IS", "R"}
    if not required.issubset(set(wb.sheetnames)):
        raise ValueError("This profile requires an integrated 14-tab-style LBO workbook.")
    for _, _, sheet, cell, label, _, _ in PROFILE:
        row = wb[sheet][cell].row
        actual = str(wb[sheet].cell(row, 2).value or "")
        if label.lower() not in actual.lower():
            raise ValueError(f"Model layout mismatch at {sheet}!B{row}: expected '{label}'.")

def inspect_template(data):
    wb = _load(data)
    _validate_layout(wb)
    entries = []
    for key, label, sheet, address, _, lo, hi in PROFILE:
        cell = wb[sheet][address]
        entries.append({
            "key": key, "input": label, "cell": f"{sheet}!{address}",
            "current": cell.value if isinstance(cell.value, (int, float)) else None,
            "proposed": cell.value if isinstance(cell.value, (int, float)) else None,
            "apply": False,
            "editable": cell.data_type != "f",
        })
    return {"sheets": wb.sheetnames, "inputs": entries}

def populate_template(data, updates):
    wb = _load(data)
    _validate_layout(wb)
    allowed = {spec[0]: spec for spec in PROFILE}
    updated = []
    if not updates:
        raise ValueError("No approved input changes were selected.")
    for key, raw in updates.items():
        if key not in allowed:
            raise ValueError(f"Unknown model input: {key}")
        _, label, sheet, address, _, lo, hi = allowed[key]
        cell = wb[sheet][address]
        if cell.data_type == "f":
            raise ValueError(f"{sheet}!{address} contains a formula and cannot be overwritten.")
        try:
            value = float(raw)
        except (TypeError, ValueError):
            raise ValueError(f"Invalid numeric value for {label}.")
        if not math.isfinite(value) or not lo <= value <= hi:
            raise ValueError(f"Out-of-range value for {label} ({lo} to {hi}).")
        if key in {"case", "tax_stepup", "tax_amort_years"}:
            if not value.is_integer():
                raise ValueError(f"{label} must be an integer.")
            value = int(value)
        cell.value = value
        updated.append(f"{sheet}!{address}")
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    out = BytesIO()
    wb.save(out)
    return out.getvalue(), updated
